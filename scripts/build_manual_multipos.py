from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs" / "Manual de Usuario MultiPOS.docx"

BLACK = "000000"
INK = "1F2925"
MUTED = "5B6761"
ACCENT = "B95231"
SOFT = "F6F3EF"
GRID = "D9D9D9"
TABLE_HEAD = "2B3732"


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def borders(cell, color=GRID):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_borders = tc_pr.first_child_found_in("w:tcBorders")
    if tc_borders is None:
        tc_borders = OxmlElement("w:tcBorders")
        tc_pr.append(tc_borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = "w:" + edge
        element = tc_borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            tc_borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "4")
        element.set(qn("w:color"), color)


def set_cell_margins(cell, top=90, start=110, bottom=90, end=110):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn("w:" + m))
        if node is None:
            node = OxmlElement("w:" + m)
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def keep_with_next(paragraph):
    p_pr = paragraph._p.get_or_add_pPr()
    keep = OxmlElement("w:keepNext")
    p_pr.append(keep)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_col_width(cell, inches):
    cell.width = Inches(inches)
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is not None:
        tc_w.set(qn("w:w"), str(int(inches * 1440)))
        tc_w.set(qn("w:type"), "dxa")


def set_font(run, size=10.5, bold=False, color=INK, italic=False):
    run.font.name = "Arial"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = RGBColor.from_string(color)


def text(p, value, bold=False, italic=False, color=INK, size=10.5):
    run = p.add_run(value)
    set_font(run, size=size, bold=bold, color=color, italic=italic)
    return run


def add_body(doc, value, first=None):
    p = doc.add_paragraph(style="Body Text")
    if first:
        text(p, first, bold=True)
    text(p, value)
    return p


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        if isinstance(item, tuple):
            text(p, item[0], bold=True)
            text(p, item[1])
        else:
            text(p, item)


def add_steps(doc, steps):
    for item in steps:
        p = doc.add_paragraph(style="List Number")
        text(p, item)


def add_heading(doc, title, level=1):
    p = doc.add_heading(title, level=level)
    for run in p.runs:
        set_font(run, size={1: 17, 2: 13, 3: 11}[level], bold=True, color=BLACK)
    p.paragraph_format.space_before = Pt(16 if level == 1 else 11)
    p.paragraph_format.space_after = Pt(6)
    keep_with_next(p)
    return p


def add_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.autofit = False
    table.style = "Table Grid"
    header = table.rows[0]
    set_repeat_table_header(header)
    for i, label in enumerate(headers):
        cell = header.cells[i]
        shade(cell, TABLE_HEAD)
        borders(cell)
        set_cell_margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        text(p, label, bold=True, color="FFFFFF", size=9.5)
        if widths:
            set_col_width(cell, widths[i])
    for idx, values in enumerate(rows):
        row = table.add_row()
        for i, value in enumerate(values):
            cell = row.cells[i]
            shade(cell, "FFFFFF" if idx % 2 == 0 else SOFT)
            borders(cell)
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if widths:
                set_col_width(cell, widths[i])
            p = cell.paragraphs[0]
            text(p, str(value), size=9.5, color=INK)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return table


def add_divider_space(doc, points=6):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(points)


def add_screenshot(doc, filename, caption):
    path = ROOT / "docs" / "screenshots" / filename
    if not path.exists():
        raise FileNotFoundError(f"Falta la captura requerida: {path}")
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(5)
    p.paragraph_format.space_after = Pt(3)
    picture = p.add_run().add_picture(str(path), width=Inches(6.4))
    picture._inline.docPr.set("descr", caption)
    picture._inline.docPr.set("title", caption)
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.space_after = Pt(9)
    text(cap, caption, size=9, italic=True, color=MUTED)


def setup_styles(doc):
    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.13
    body = doc.styles["Body Text"]
    body.base_style = normal
    body.paragraph_format.space_after = Pt(7)
    body.paragraph_format.line_spacing = 1.13
    for name, size in (("Title", 26), ("Subtitle", 13), ("Heading 1", 17), ("Heading 2", 13), ("Heading 3", 11)):
        style = doc.styles[name]
        style.font.name = "Arial"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
        style.font.size = Pt(size)
        style.font.bold = name != "Subtitle"
        style.font.color.rgb = RGBColor.from_string(BLACK)
    for name in ("List Bullet", "List Number"):
        style = doc.styles[name]
        style.font.name = "Arial"
        style.font.size = Pt(10.5)


def add_footer(section):
    p = section.footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(3)
    text(p, "Manual de Usuario MultiPOS  |  Operación diaria", size=8.5, color=MUTED)


def cover(doc):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(122)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    text(p, "MultiPOS", size=13, bold=True, color=ACCENT)
    p = doc.add_paragraph(style="Title")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    text(p, "Manual de Usuario", size=27, bold=True, color=BLACK)
    p = doc.add_paragraph(style="Subtitle")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    text(p, "Guía completa para operar ventas, inventario, caja y administración", size=13, color=MUTED)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(40)
    text(p, "Versión operativa", size=10.5, bold=True, color=INK)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    text(p, "Preparado para el equipo de caja, ventas, inventario y administración", size=10.5, color=MUTED)
    doc.add_page_break()


def build():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.68)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.72)
    section.right_margin = Inches(0.72)
    setup_styles(doc)
    add_footer(section)
    cover(doc)

    add_heading(doc, "Propósito y alcance", 1)
    add_body(doc, "Este manual explica cómo usar MultiPOS para la operación diaria de un negocio: facturar, controlar inventario, registrar compras, manejar caja, administrar créditos, generar reportes y mantener la configuración de la empresa. Está escrito para usuarios operativos y administradores.")
    add_body(doc, "La pantalla y las opciones que ve cada persona dependen de su rol y permisos. Si una opción no aparece en el menú, solicite al administrador que revise su acceso; no intente resolverlo compartiendo contraseñas.")

    add_heading(doc, "Cómo leer este manual", 2)
    add_bullets(doc, [
        ("Ruta de menú. ", "Indica dónde se encuentra una pantalla, por ejemplo: Administración > Caja."),
        ("Acción principal. ", "Describe la secuencia recomendada para realizar una tarea."),
        ("Importante. ", "Aclara una regla de negocio, un efecto en el inventario o una condición de permisos."),
    ])

    add_heading(doc, "Contenido", 1)
    toc = [
        "1. Acceso, sesión y navegación", "2. Inicio y alertas", "3. Catálogo de productos e inventario",
        "4. Clientes y proveedores", "5. Venta estándar", "6. Venta táctil", "7. Lista de ventas, tickets y comprobantes",
        "8. Cotizaciones y preventas", "9. Compras y pedidos a proveedor", "10. Traslados y auditoría de inventario",
        "11. Caja y diario digital", "12. Cuentas por cobrar y por pagar", "13. Gastos, empleados y planillas",
        "14. Reportes", "15. Empresa, configuración e impresión", "16. Usuarios y permisos",
        "17. Sincronización, modo local y respaldos", "18. Rutinas recomendadas y solución de problemas",
    ]
    for item in toc:
        p = doc.add_paragraph(style="List Bullet")
        text(p, item)

    add_heading(doc, "1. Acceso sesión y navegación", 1)
    add_heading(doc, "Iniciar sesión", 2)
    add_body(doc, "En la pantalla de acceso seleccione la empresa, escriba su usuario y contraseña, y pulse Iniciar sesión. La empresa elegida define los datos, documentos e inventario que podrá consultar durante esa sesión.")
    add_steps(doc, [
        "Seleccione la empresa correcta en la lista Empresa.",
        "Escriba su Usuario y Contraseña.",
        "Pulse Iniciar sesión. Si el botón permanece desactivado, revise que los tres campos estén completos.",
    ])
    add_body(doc, "El botón Ajustes avanzados sirve para indicar la URL del API cuando el navegador debe conectarse a otro servidor. Es una configuración técnica; úsela solamente con la dirección proporcionada por el administrador.", first="Importante. ")
    add_heading(doc, "Navegación principal", 2)
    add_body(doc, "En escritorio, el menú lateral permanece disponible y se puede contraer. En móvil se abre desde el botón de menú de la cabecera. La acción Nueva venta está siempre al inicio para que la operación de caja sea rápida.")
    add_table(doc, ["Grupo", "Pantallas principales"], [
        ["Uso diario", "Inicio, Ventas, Productos, Caja y Nueva venta"],
        ["Facturación", "Venta táctil, Cotizaciones, PreVentas, Cuentas por cobrar, Compras, Pedidos a proveedor y Traslados"],
        ["Directorio", "Clientes y Proveedores"],
        ["Administración", "Cuentas por pagar, Gastos, Reportes, Auditoría, Empleados y Planillas"],
        ["Empresa", "Información de empresa, Usuarios y permisos, Configuración y Ayuda"],
    ], [1.25, 5.45])
    add_body(doc, "La cabecera muestra la empresa, sucursal o caja activa cuando corresponde, el estado de sincronización y el menú del usuario. Desde ese menú puede bloquear la pantalla o cerrar sesión.")

    add_heading(doc, "2. Inicio y alertas", 1)
    add_body(doc, "Inicio es el tablero de la operación. Resume las ventas de hoy, el estado de caja y los gastos del día. También muestra una tendencia de ventas de los últimos 14 días, las ventas recientes y pendientes que requieren atención.")
    add_bullets(doc, [
        ("Ventas de hoy. ", "Monto acumulado y número de documentos del día."),
        ("Caja. ", "Indica si el turno actual está abierto y desde cuándo."),
        ("Gastos de hoy. ", "Total registrado para la fecha actual."),
        ("Requiere atención. ", "Avisa principalmente sobre productos agotados o bajo mínimo."),
    ])
    add_body(doc, "Use el tablero para detectar una excepción y vaya al módulo correspondiente. No sustituye el cierre de caja ni los reportes detallados.")
    add_screenshot(doc, "manual-dashboard.png", "Pantalla Inicio: resumen operativo, tendencia, ventas recientes y alertas.")

    add_heading(doc, "3. Catálogo de productos e inventario", 1)
    add_heading(doc, "Consultar y filtrar productos", 2)
    add_body(doc, "Ruta de menú: Productos. Use el campo de búsqueda para localizar por nombre, SKU, código de barras o código rápido. También puede filtrar por existencia, proveedor y rango de vencimiento. Las filas muestran existencia, precio, categoría, ubicación, ISV, vencimiento, tipo y acceso al historial de movimientos.")
    add_bullets(doc, [
        ("Con existencia, sin existencia y bajo mínimo. ", "Ayudan a priorizar reposición."),
        ("Historial. ", "Expone las entradas, salidas, ajustes y referencias que explican una variación de inventario."),
        ("Edición. ", "Permite corregir la ficha del producto; el acceso puede depender del rol."),
    ])
    add_screenshot(doc, "manual-products.png", "Pantalla Productos: filtros, alta de producto y tabla del catálogo.")
    add_heading(doc, "Crear o editar un producto", 2)
    add_steps(doc, [
        "Pulse Nuevo producto o Edición en la fila correspondiente.",
        "En la pestaña Producto complete nombre, unidad, proveedor, costo, ISV, precio principal, existencia mínima, categoría y ubicación. Agregue marca, imagen, lote y fecha de vencimiento cuando correspondan.",
        "Defina el tipo: Producto, Servicio, Insumo o Kit. Guarde los cambios.",
    ])
    add_table(doc, ["Tipo", "Comportamiento"], [
        ["Producto", "Se vende en POS y controla existencia."],
        ["Servicio", "Se factura, pero no descuenta inventario."],
        ["Insumo", "Se usa como insumo y no aparece en la venta POS."],
        ["Kit o combo", "Se vende como una línea, pero descuenta la existencia de cada componente tipo Producto."],
    ], [1.45, 5.25])
    add_body(doc, "Para un Kit, abra la pestaña Combo, busque productos tipo Producto y agregue la cantidad que contiene cada uno. El stock del kit no se captura manualmente; se determina por sus componentes.", first="Kit. ")
    add_heading(doc, "Precios y volumen", 2)
    add_body(doc, "Cada producto puede tener cuatro listas de precio. La venta selecciona la lista del documento; además, puede definir tramos por cantidad. Cuando una cantidad alcanza un tramo, se utiliza el precio del mayor mínimo alcanzado; si no, se aplica la lista seleccionada en la venta.")
    add_body(doc, "El precio de venta mostrado al cliente ya incorpora la lógica fiscal definida para el producto. En el resumen de venta, el ISV se muestra desglosado como parte del precio final, no como un cargo adicional inesperado.", first="Precio con impuesto incluido. ")
    add_heading(doc, "Existencia por ubicación", 2)
    add_body(doc, "Cuando hay ubicaciones configuradas, un administrador puede consultar el desglose de existencia por bodega dentro de la ficha del producto. Para distribuir existencias entre ubicaciones use Traslados; para corregir diferencias use Auditoría de inventario.")

    add_heading(doc, "4. Clientes y proveedores", 1)
    add_heading(doc, "Clientes", 2)
    add_body(doc, "Ruta de menú: Directorio > Clientes. Busque por nombre, teléfono, RTN o código. Desde el directorio puede crear y editar clientes con nombre, dirección, teléfono, RTN, notas y una lista de precio predeterminada.")
    add_body(doc, "Durante una venta, escriba directamente en Cliente. El sistema busca coincidencias mientras escribe. Si el nombre es nuevo, puede completar dirección, teléfono y RTN y aceptar guardar el cliente para reutilizarlo después.")
    add_body(doc, "Para vender a crédito, seleccione o registre un cliente identificable. Consumidor final es útil para venta inmediata, pero no debe reemplazar los datos de una cuenta por cobrar.", first="Crédito. ")
    add_screenshot(doc, "manual-customers.png", "Pantalla Clientes: directorio, búsqueda y acceso para crear o editar registros.")
    add_heading(doc, "Proveedores", 2)
    add_body(doc, "Ruta de menú: Directorio > Proveedores. Registre nombre, teléfono, correo, RTN y dirección. El proveedor se puede asociar a productos, compras y pedidos para mantener la trazabilidad de abastecimiento.")

    add_heading(doc, "5. Venta estándar", 1)
    add_body(doc, "Ruta de menú: Nueva venta. Es el flujo optimizado para teclado, lector de códigos y una grilla de líneas. Complete los datos del documento y agregue productos antes de cobrar.")
    add_heading(doc, "Flujo recomendado de una venta", 2)
    add_steps(doc, [
        "Verifique fecha, términos y lista de precios. Use Contado, Tarjeta, Efectivo o Crédito según aplique.",
        "Busque al cliente en el campo Cliente. Para una venta nueva puede escribirlo y guardar sus datos si desea conservarlos.",
        "Escanee el código, escriba código rápido o SKU en Agregar productos. También puede abrir Catálogo para buscar por nombre, proveedor o categoría.",
        "Revise cantidad, unidad, precio final, descuento autorizado, ISV y total de cada línea.",
        "Ingrese el monto cobrado cuando corresponda. Revise total, pago, cambio o saldo.",
        "Guarde la venta. Puede imprimir ticket o abrir el comprobante después de guardarla.",
    ])
    add_screenshot(doc, "manual-new-sale.png", "Pantalla Nueva venta: datos del cliente, condiciones, lista de precios y acciones rápidas.")
    add_heading(doc, "Reglas de precio y descuentos", 2)
    add_body(doc, "El sistema obtiene el precio del catálogo y la lista de precios elegida. Modificar un precio manualmente o aplicar descuento requiere el permiso sales.price_override. Si no lo tiene, esos campos quedan en solo lectura; solicite al administrador una autorización en lugar de buscar una forma de eludir la regla.")
    add_body(doc, "Los precios y el total utilizan el ISV configurado por producto. El resumen muestra Subtotal, ISV incluido y Total para que pueda explicarlo al cliente con claridad.")
    add_heading(doc, "Atajos operativos", 2)
    add_table(doc, ["Acción", "Atajo"], [
        ["Guardar venta", "F5 o Ctrl + Enter"],
        ["Guardar e imprimir ticket", "F8 o Ctrl + Shift + Enter"],
        ["Buscar cliente", "F2"],
        ["Buscar productos", "F4 o Ctrl + K"],
        ["Enfocar código o barras", "F9"],
        ["Eliminar línea seleccionada o última", "F10"],
    ], [2.8, 3.9])
    add_body(doc, "Algunos navegadores reservan teclas de función, por ejemplo F5 para recargar. Si un atajo no responde, use el botón visible equivalente o la combinación Ctrl indicada. La aplicación intercepta los atajos cuando el navegador lo permite.")
    add_heading(doc, "Venta a crédito", 2)
    add_body(doc, "Seleccione un término de crédito y establezca el pago inicial si existe. Al guardar, la venta genera un saldo en Cuentas por cobrar. El saldo se reduce posteriormente mediante abonos registrados en ese módulo.")

    add_heading(doc, "6. Venta táctil", 1)
    add_body(doc, "Ruta de menú: Facturación > Venta táctil. Está diseñada para pantallas táctiles y operación rápida. Muestra tarjetas de producto, búsqueda por nombre, SKU o código, carrito lateral en escritorio y resumen de cobro en móvil.")
    add_steps(doc, [
        "Elija o busque un producto y tóquelo para agregarlo al carrito.",
        "Aumente o reduzca cantidades desde el carrito; quite una línea si fue agregada por error.",
        "Seleccione cliente, término y lista de precios. El cliente es obligatorio para crédito.",
        "Ingrese el importe recibido y confirme el cobro. El sistema calcula cambio o saldo.",
    ])
    add_body(doc, "Puede marcar productos como favoritos para que aparezcan antes en esta pantalla. Los favoritos se guardan para la organización y ayudan a acelerar artículos de alta rotación.")
    add_screenshot(doc, "manual-touch-sale.png", "Pantalla Venta táctil: favoritos, tarjetas de producto y carrito de cobro.")

    add_heading(doc, "7. Lista de ventas tickets y comprobantes", 1)
    add_heading(doc, "Consultar ventas", 2)
    add_body(doc, "Ruta de menú: Ventas. Use filtros por texto, cliente, tipo de término, fecha y estado. La tabla puede mostrar número de factura, fecha, cliente, términos, total, pago, saldo, hora, estado y vendedor. Las columnas visibles se pueden ajustar desde la propia pantalla.")
    add_bullets(doc, [
        ("Comprobante. ", "Abre el documento de venta en formato carta, desde donde se puede imprimir o descargar PDF."),
        ("Ticket. ", "Abre el recibo preparado para impresora térmica de 80 mm."),
        ("Edición. ", "La edición de una venta es exclusiva de administración."),
        ("Eliminación. ", "Requiere el permiso sales.delete y motivo. La venta deja de contar en reportes y caja, pero conserva bitácora de quién la eliminó."),
    ])
    add_heading(doc, "Imprimir documentos", 2)
    add_body(doc, "En un ticket o comprobante use Imprimir del navegador y seleccione la impresora correcta. Desde el comprobante puede usar Imprimir o PDF del navegador, o Descargar PDF del servidor. El ticket y el comprobante utilizan la información de empresa y la configuración de factura activa.")

    add_heading(doc, "8. Cotizaciones y preventas", 1)
    add_body(doc, "Ruta de menú: Facturación > Cotizaciones o PreVentas. Ambas registran un documento antes de convertirlo en venta. No descuentan inventario mientras sigan como cotización o preventa.")
    add_steps(doc, [
        "Cree una nueva cotización o preventa.",
        "Seleccione cliente si aplica, agregue productos y cantidades, y complete referencia de servicio, mesa o para llevar si se usa en el negocio.",
        "Guarde el documento. Si se habilitó cocina o preparación, puede imprimir la orden desde el navegador.",
        "Cuando el cliente confirme, use Convertir a venta. Revise el resultado y continúe con el cobro como una venta normal.",
    ])
    add_body(doc, "Los productos pueden marcarse para incluirse en la orden de cocina o bodega. Esto es útil para separar lo que debe preparar el equipo de lo que solo aparece como referencia comercial.")

    add_heading(doc, "9. Compras y pedidos a proveedor", 1)
    add_heading(doc, "Registrar una compra", 2)
    add_body(doc, "Ruta de menú: Facturación > Compras. Registre el ingreso de mercadería seleccionando proveedor, productos, cantidades, costo e ISV. Elija Contado o Crédito. Una compra registrada aumenta la existencia de los productos; si queda a crédito, genera un saldo en Cuentas por pagar.")
    add_steps(doc, [
        "Seleccione el proveedor o cree uno antes de registrar la compra.",
        "Busque y agregue productos. Capture cantidad, costo e ISV de cada línea.",
        "Elija la condición de pago y guarde. Revise el total y pendiente en el listado.",
    ])
    add_screenshot(doc, "manual-purchases.png", "Pantalla Compras: historial, total comprado y registro de mercadería.")
    add_heading(doc, "Pedidos a proveedor", 2)
    add_body(doc, "Ruta de menú: Facturación > Pedidos a proveedor. Los pedidos sirven para planificar abastecimiento sin aumentar existencias. Seleccione proveedor, fecha estimada, productos, cantidades, precio de referencia y notas. Después actualice el estado conforme se solicite, envíe, reciba o cancele el pedido.")

    add_heading(doc, "10. Traslados y auditoría de inventario", 1)
    add_heading(doc, "Traslados", 2)
    add_body(doc, "Ruta de menú: Facturación > Traslados. Cree ubicaciones o bodegas cuando sea necesario y elabore un traslado entre origen y destino. Agregue productos y cantidades, guarde el borrador, luego envíelo y recíbalo al llegar. El listado permite consultar estado, origen, destino y contenido.")
    add_bullets(doc, [
        ("Borrador. ", "Todavía puede corregirse; no representa una recepción final."),
        ("Enviar. ", "Marca el traslado en tránsito."),
        ("Recibir. ", "Confirma la llegada a la ubicación de destino."),
        ("Importar y exportar JSON. ", "Permite intercambiar un borrador de traslado cuando el procedimiento operativo lo requiera."),
    ])
    add_heading(doc, "Auditoría de inventario", 2)
    add_body(doc, "Ruta de menú: Administración > Auditoría inventario. Use este módulo para registrar un ajuste con motivo, notas y líneas. Escriba un cambio positivo para sumar existencia o negativo para restar. Cada ajuste queda en la lista reciente con número, fecha, usuario y detalle.")
    add_body(doc, "No use una venta o una compra ficticia para corregir inventario. Eso distorsiona sus reportes y cuentas; use un ajuste auditado.", first="Importante. ")

    add_heading(doc, "11. Caja y diario digital", 1)
    add_body(doc, "Ruta de menú: Caja. La caja se maneja por turno. Desde esta pantalla puede abrir turno, consultar la actividad, registrar movimientos de efectivo y cerrar con arqueo.")
    add_heading(doc, "Abrir turno", 2)
    add_steps(doc, [
        "Compruebe que no tenga un turno abierto anterior.",
        "Ingrese el fondo inicial. Si la caja inicia vacía, escriba 0.",
        "Pulse Abrir turno. Desde ese momento las ventas y gastos del usuario dentro del intervalo aparecen en el diario.",
    ])
    add_heading(doc, "Durante el turno", 2)
    add_body(doc, "El resumen muestra efectivo esperado, ventas del turno y cobros con tarjeta. Los pagos con tarjeta no se consideran efectivo en el cajón. Registre retiros, ingresos y ajustes manuales con monto y nota; por ejemplo, un retiro para depósito.")
    add_heading(doc, "Cerrar turno", 2)
    add_steps(doc, [
        "Cuente físicamente el efectivo del cajón.",
        "Ingrese el efectivo contado y una nota si hay diferencia.",
        "Compare contra el efectivo esperado y cierre el turno.",
        "Imprima el cierre si necesita conservarlo con la documentación física.",
    ])
    add_body(doc, "El efectivo sugerido usa ventas del usuario dentro de la sesión, gastos del mismo usuario y movimientos manuales. Un abono recibido hoy de una factura de crédito emitida en otro turno no entra automáticamente en ese cálculo; si se recibió en efectivo, regístrelo o considérelo en el arqueo.", first="Cómo interpretar la diferencia. ")
    add_body(doc, "La vista de historial permite consultar un día o cajero anterior sin abrir una caja. Sirve para revisar cierres, ventas y movimientos de un turno ya terminado.")
    add_screenshot(doc, "manual-cash.png", "Pantalla Caja: turno abierto, efectivo esperado, actividad y acciones de cierre.")

    add_heading(doc, "12. Cuentas por cobrar y por pagar", 1)
    add_heading(doc, "Cuentas por cobrar", 2)
    add_body(doc, "Ruta de menú: Facturación > Cuentas por cobrar. Muestra facturas de crédito con saldo, cliente, fecha de vencimiento y estado. Busque por cliente o número de factura, filtre vencidas si lo necesita y registre pagos o recargos.")
    add_steps(doc, [
        "Localice la factura con saldo pendiente.",
        "Seleccione Pago, escriba el monto y confirme el abono.",
        "Si existe interés o mora, use Recargo, capture monto y nota. Revise el nuevo saldo.",
    ])
    add_screenshot(doc, "manual-receivables.png", "Pantalla Cuentas por cobrar: saldos, búsqueda y registro de abonos.")
    add_heading(doc, "Cuentas por pagar", 2)
    add_body(doc, "Ruta de menú: Administración > Cuentas por pagar. Funciona de forma equivalente para compras a crédito: permite consultar referencia, proveedor, total, recargos, pago realizado y saldo. Registre pagos y recargos sobre la compra correcta.")
    add_body(doc, "Antes de registrar un abono, confirme documento, proveedor o cliente y monto. Los abonos afectan el saldo; un dato equivocado complica el arqueo y los estados de cuenta.")

    add_heading(doc, "13. Gastos empleados y planillas", 1)
    add_heading(doc, "Gastos", 2)
    add_body(doc, "Ruta de menú: Administración > Gastos. Los administradores pueden crear libros de gastos y sus categorías, luego registrar fecha, monto, categoría y notas. El listado se filtra por fecha, libro o categoría y muestra quién registró cada gasto.")
    add_body(doc, "Use un libro y categoría consistentes, por ejemplo Operaciones > Combustible. Esto hace que el reporte de gastos sea útil y que el diario de caja se explique mejor.")
    add_heading(doc, "Empleados y planillas", 2)
    add_body(doc, "Ruta de menú: Administración > Empleados y Planillas. Empleados guarda código, nombre, documento, contacto, puesto, fecha de ingreso, estado y notas. Planillas permite crear periodos mensuales, agregar líneas por empleado, registrar deducciones y cerrar el periodo.")
    add_body(doc, "Los usuarios con permiso de consulta pueden revisar información autorizada, mientras que crear o cerrar planilla corresponde a administración. Verifique bruto, deducciones y neto antes de cerrar.")

    add_heading(doc, "14. Reportes", 1)
    add_body(doc, "Ruta de menú: Administración > Reportes. Utilice fechas y filtros para analizar operación. Los permisos determinan cuáles pestañas y datos puede consultar.")
    add_table(doc, ["Reporte", "Qué responde"], [
        ["Ventas", "Cantidad de documentos, total facturado y exportación de resultados a CSV."],
        ["Inventario", "Existencia, costo, precio y valor estimado del inventario."],
        ["Productos más vendidos", "Cantidades e importes por producto en el periodo."],
        ["Gastos", "Registros y totales por categoría."],
        ["Planillas", "Periodos, empleados, bruto, deducciones y neto."],
    ], [2.2, 4.5])
    add_body(doc, "Defina primero el rango de fechas y luego actualice el reporte. Al exportar ventas, verifique que los filtros activos sean los que desea entregar.")
    add_screenshot(doc, "manual-reports.png", "Pantalla Reportes: pestañas por tipo, fechas y resumen del periodo.")

    add_heading(doc, "15. Empresa configuración e impresión", 1)
    add_heading(doc, "Información de empresa", 2)
    add_body(doc, "Ruta de menú: Empresa > Información empresa. Mantenga actualizados nombre comercial, RTN, teléfonos, correo, dirección fiscal, ciudad, departamento, sitio web, eslogan, logo, fecha, idioma y moneda. Estos datos se usan en documentos y comprobantes.")
    add_body(doc, "Si usa un logo, introduzca una URL HTTPS o una imagen PNG/JPEG compatible. Verifique un comprobante de prueba después de cambiarlo, porque el servidor debe poder obtener la imagen para incluirla en el PDF.")
    add_heading(doc, "Configuración", 2)
    add_body(doc, "Ruta de menú: Empresa > Configuración. Solo administración. Agrupa ajustes de apariencia, flujo de ventas, comportamiento del POS, ticket, comprobante, preventa o cocina, configuración SAR, conexión, sincronización, respaldo e importaciones.")
    add_table(doc, ["Sección", "Uso"], [
        ["Apariencia", "Elija tema visual, incluido el tema nocturno, y ordene Facturación con énfasis mixto, preventa o caja."],
        ["Punto de venta", "Ajusta el comportamiento que consumen venta estándar y táctil."],
        ["Factura y ticket", "Define encabezado, pie, formato de comprobante y opciones de impresión."],
        ["SAR Honduras", "Configura serie y texto legal que aparecerá en la facturación configurada."],
        ["Conexión y nube", "Configura base del API y consulta o ejecuta sincronización cuando esté habilitada."],
        ["Respaldo e importación", "Exporta un respaldo JSON e importa plantillas de productos, clientes o proveedores según el procedimiento autorizado."],
    ], [1.55, 5.15])
    add_body(doc, "Cambios de configuración afectan a más de un usuario. Realice una prueba controlada, en especial antes de cambiar configuración fiscal, impresión, listas de precios o conexión.")
    add_screenshot(doc, "manual-settings.png", "Pantalla Configuración: temas, flujo de ventas y ajustes operativos.")

    add_heading(doc, "16. Usuarios y permisos", 1)
    add_body(doc, "Ruta de menú: Empresa > Usuarios y permisos. El administrador crea usuarios, cambia nombre, rol, contraseña y estado activo. También puede aplicar una matriz fina de permisos para conceder o restringir acciones concretas.")
    add_table(doc, ["Rol", "Alcance base"], [
        ["Administrador", "Acceso completo. La matriz fina no le restringe permisos."],
        ["Vendedor", "Acceso operativo con consulta de reportes, traslados, cuentas por cobrar y pagar, y registro de compras según la base actual."],
        ["Cajero", "Operación de caja y accesos base a cuentas por cobrar, cuentas por pagar y compras; el administrador puede ajustar cada caso."],
    ], [1.5, 5.2])
    add_body(doc, "Entre los permisos ajustables están: ver reportes, hacer traslados, cuentas por cobrar, cuentas por pagar, registrar o ver compras, ver gastos, ver planillas, modificar precio o descuento, y eliminar ventas.")
    add_body(doc, "Cuando cambian rol o permisos, el sistema invalida sesiones antiguas para proteger los datos. El usuario puede recibir PERM_STALE y tendrá que iniciar sesión de nuevo.", first="Sesión. ")

    add_heading(doc, "17. Sincronización modo local y respaldos", 1)
    add_heading(doc, "Estado de conexión", 2)
    add_body(doc, "La cabecera indica si el equipo está trabajando localmente, conectado a la nube o si existen eventos pendientes o con error. El modo local offline está pensado para equipos Windows con API y base local; la conexión de nube facilita varias cajas, celulares y acceso remoto cuando el servicio está configurado.")
    add_body(doc, "Antes de una operación crítica en varias cajas, revise el estado de sincronización. Si aparecen pendientes o errores, no asuma que el otro equipo ya ve su documento.")
    add_heading(doc, "Respaldo", 2)
    add_steps(doc, [
        "Entre a Empresa > Configuración > Respaldo como administrador.",
        "Descargue el archivo JSON y guárdelo fuera del equipo de caja, con fecha en el nombre.",
        "Conserve varias copias según la política del negocio y proteja el archivo, porque contiene información operativa.",
    ])
    add_body(doc, "La restauración completa no es una acción cotidiana ni debe hacerse sobre datos de producción sin validar el archivo y el alcance. Siga el procedimiento interno o pida apoyo técnico antes de importar o reemplazar datos.")

    add_heading(doc, "18. Rutinas recomendadas y solución de problemas", 1)
    add_heading(doc, "Rutina de apertura", 2)
    add_steps(doc, [
        "Inicie sesión en la empresa correcta.",
        "Revise Inicio para detectar productos agotados o pendientes.",
        "Abra Caja con el fondo físico inicial y confirme que la fecha sea correcta.",
        "Verifique impresora, lector y conexión antes de atender la primera venta.",
    ])
    add_heading(doc, "Rutina durante el día", 2)
    add_bullets(doc, [
        "Registre cada venta, compra, gasto y movimiento en su módulo real; evite documentos ficticios.",
        "Use el cliente real para ventas a crédito y capture correctamente el pago inicial.",
        "Revise alertas de bajo stock y genere pedidos antes de agotar artículos críticos.",
        "Registre movimientos manuales de caja en el momento en que suceden.",
    ])
    add_heading(doc, "Rutina de cierre", 2)
    add_steps(doc, [
        "Revise que las ventas y pagos del día estén guardados.",
        "Cuente efectivo, registre diferencia explicada y cierre Caja.",
        "Imprima o guarde el cierre requerido por el negocio.",
        "Ejecute el respaldo diario o confirme que la sincronización esté al día.",
    ])
    add_heading(doc, "Problemas frecuentes", 2)
    add_table(doc, ["Situación", "Qué hacer"], [
        ["No aparece una opción", "Revise el rol y permisos con un administrador. La aplicación oculta funciones no autorizadas."],
        ["No se puede cambiar precio o descuento", "Solicite sales.price_override; el precio se protege desde el servidor."],
        ["F5 o una tecla de función recarga el navegador", "Use el botón visible o el atajo alternativo Ctrl indicado en Venta estándar."],
        ["La pantalla cerró sesión con PERM_STALE", "Inicie sesión otra vez; un administrador cambió rol o permisos."],
        ["El ticket no imprime", "Compruebe que la impresora esté instalada en el sistema operativo y selecciónela en el diálogo de impresión."],
        ["La caja no cuadra", "Revise ventas, gastos y movimientos del turno; compare efectivo físico y deje una nota de diferencia."],
        ["No hay conexión", "Revise la URL del API y el indicador de sincronización. No cambie ajustes avanzados sin la dirección autorizada."],
    ], [2.25, 4.45])
    add_body(doc, "Para ayuda rápida dentro de la aplicación use Empresa > Ayuda y FAQ. También existe Ayuda pública desde la pantalla de inicio de sesión para consultar información sin entrar al sistema.")

    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
